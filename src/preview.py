"""Print a JSON preview of one file: spreadsheet values, or an email's plain text."""
import csv, email, json, sys
from email import policy
from pathlib import Path

ROWS, COLS = 2000, 40
path = Path(sys.argv[1])
ext = path.suffix.lower()


def sheet(name, rows):
    out = []
    for r in rows:
        if len(out) == ROWS:
            return {'name': name, 'rows': out, 'truncated': True}
        out.append(list(r)[:COLS])
    return {'name': name, 'rows': out, 'truncated': False}


if ext in ('.xlsx', '.xlsm'):
    from openpyxl import load_workbook
    values = load_workbook(path, read_only=True, data_only=True)
    formulas = load_workbook(path, read_only=True)

    def cells(vs, fs):
        for vrow, frow in zip(vs.iter_rows(values_only=True), fs.iter_rows(values_only=True)):
            yield [v if v is not None or not str(f).startswith('=') else {'f': f} for v, f in zip(vrow, frow)]

    out = {'sheets': [sheet(vs.title, cells(vs, fs)) for vs, fs in zip(values.worksheets, formulas.worksheets)]}
elif ext == '.xls':
    import xlrd
    book = xlrd.open_workbook(path)
    out = {'sheets': [sheet(s.name, (s.row_values(i) for i in range(s.nrows))) for s in book.sheets()]}
elif ext == '.csv':
    with open(path, newline='', encoding='utf-8-sig', errors='replace') as f:
        out = {'sheets': [sheet(path.name, csv.reader(f))]}
elif ext == '.eml':
    m = email.message_from_bytes(path.read_bytes(), policy=policy.default)
    body = m.get_body(preferencelist=('plain',))
    out = {'email': {
        'headers': {k: str(m[k]) for k in ('From', 'To', 'Cc', 'Date', 'Subject') if m[k]},
        'text': body.get_content() if body else '',
        'attachments': [a.get_filename() for a in m.iter_attachments()]}}
else:
    out = {}
print(json.dumps(out, default=str))
