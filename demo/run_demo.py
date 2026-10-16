"""Drive the made-up weekly borrowing base series through a running server, then score each forward week against truth.json.

Usage: .venv/bin/python demo/run_demo.py demo/out http://127.0.0.1:8780 DATA_ROOT [first-label]
Each week is staged only when the agent may see it, so it never sees the future. The sources the agent fetches from 08-16
live in DATA_ROOT-sources: warehouse.sqlite and mailbox/. Scoring and the checks are test tooling, not part of the system.
"""
import csv, json, shutil, sqlite3, sys, time, urllib.parse, urllib.request
from pathlib import Path
from openpyxl import load_workbook

SRC = Path(sys.argv[1])
SERIES = 'Weekly borrowing base'
API = sys.argv[2] + '/api/larkfield/' + urllib.parse.quote(SERIES)
DATA = Path(sys.argv[3]) / 'larkfield' / SERIES
SOURCES = Path(sys.argv[3] + '-sources').resolve()
WAREHOUSE, MAILBOX = SOURCES / 'warehouse.sqlite', SOURCES / 'mailbox'
EXAMPLES = ['2026-07-05', '2026-07-12', '2026-07-19', '2026-07-26']
ORDER = EXAMPLES + ['2026-08-02', '2026-08-09', '2026-08-16', '2026-08-23', '2026-08-30']
ECB = 'https://data-api.ecb.europa.eu/service/data/EXR/D.USD.EUR.SP00.A?startPeriod=<date>&endPeriod=<date>&format=csvdata'
SOURCES_STEER = f"""From this week on we stop uploading two files; fetch them yourself:
- Inventory: the warehouse database {WAREHOUSE} (SQLite), table inventory_snapshot, one snapshot per as_of date. Use the latest snapshot on or before the week end.
- Loan activity: Granite Point emails it each Monday to the mailbox folder {MAILBOX} (.eml files with the xlsx attached).
Also new: Maison Verte SARL is billed in EUR from this week (see the Currency column of the AR aging). Convert it at the ECB euro reference rate for the week end, or the last business day before it, from the ECB data API: {ECB}"""


def call(method, path='', data=b'', **q):
    url = API + path + ('?' + urllib.parse.urlencode(q) if q else '')
    with urllib.request.urlopen(urllib.request.Request(url, data=data if method == 'POST' else None, method=method)) as r:
        body = r.read().decode()
    return json.loads(body) if body[:1] in '{[' else body


def state(label):
    return call('GET', '/' + label)


def upload(label, path, pane, by='alex', onto=None):
    q = {'by': by, 'pane': pane, 'name': path.name, **({'onto': onto} if onto else {})}
    call('POST', f'/{label}/upload', path.read_bytes(), **q)
    print(f'  {by} uploaded {path.name} ({pane}{", onto " + onto if onto else ""})')


def steer(label, text, by='alex'):
    call('POST', f'/{label}/steer', json.dumps({'text': text}).encode(), by=by)
    print(f'  {by} steered: {text.splitlines()[0]}')


def wait(label):
    """Wait for the runs started by our last events to finish."""
    mark = state(label)['log'][-1]['seq']
    while True:
        time.sleep(10)
        s = state(label)
        if s['run'] is None and any(e['op'] == 'done' and e['seq'] > mark for e in s['log']):
            for e in s['log']:
                if e['seq'] > mark and e['op'] == 'done':
                    print(f"  run {e['run']} done{' (aborted)' if e['aborted'] else ''} ${e['cost']:.2f}: {e['summary']}")
            return s


def started(label):
    while state(label)['run'] is None:
        time.sleep(3)


def screen(s):
    for i in s['doc']['items']:
        print(f"  [{i['kind']}] {i['id']}: {i['file']}{'  | ' + i['headline'] if i.get('headline') else ''}{'  | source: ' + i['source'] if i.get('source') else ''}")
        for n in i.get('notes') or []:
            q = f" {n['options']} default {n.get('default', 0)}" if n.get('options') else ''
            f = f" files {n['files']}" if n.get('files') else ''
            print(f"      - {n.get('ref') or ''} {n['text']}{q}{f}")


def lines(path):
    ws = load_workbook(path, data_only=True)['Base']
    out = []
    for row in ws.iter_rows(values_only=True):
        label = next((c for c in row if isinstance(c, str)), None)
        nums = [c for c in row if isinstance(c, (int, float)) and not isinstance(c, bool)]
        if label and nums:
            out.append((label.strip().lower(), nums[0]))
    return out


def score(label, s, key='lines'):
    truth = json.loads((SRC / label / 'truth.json').read_text())[key]
    outs = [i for i in s['doc']['items'] if i['kind'] == 'output' and i['file'] and i['file'].startswith('working/')]
    if not outs:
        print('  SCORE: no agent output on screen')
        return
    got = lines(DATA / label / outs[-1]['file'])
    used, ok, bad = set(), 0, []
    for name, value in truth:
        match = next((i for i, (n, v) in enumerate(got) if n == name.lower() and i not in used), None)
        if match is not None:
            used.add(match)
        if match is not None and abs(got[match][1] - value) <= (0.0001 if abs(value) < 1 else 1):
            ok += 1
        else:
            bad.append((name, value, got[match][1] if match is not None else None))
    print(f'  SCORE {label} ({outs[-1]["file"]}): {ok}/{len(truth)} lines within $1')
    for name, want, have in bad:
        print(f'      {name}: want {want:,.2f}, got {have if have is None else f"{have:,.2f}"}')


def check(label, s):
    items = s['doc']['items']
    placeholders = [i['id'] for i in items if i['id'][:1] == 'u' and i['id'][1:].split('_')[0].isdigit()]
    no_headline = [i['id'] for i in items if i['kind'] == 'output' and i['file'] and not i.get('headline')]
    sourced = [i['id'] for i in items if i.get('source')]
    cited = sorted({f for i in items for n in i.get('notes') or [] for f in n.get('files') or []})
    process = DATA / label / 'working' / 'process.md'
    print(f"  CHECK {label}: from {s['log'][0].get('from')}; placeholder ids {placeholders}; outputs without headline {no_headline}; "
          f"sourced {sourced}; cited {cited}; process.md {process.stat().st_size if process.exists() else 0} bytes")


def stage_sources(label):
    """Add this week's warehouse snapshot; the mailbox gets its emails separately."""
    rows = list(csv.DictReader(open(SRC / label / 'sources' / 'inventory.csv')))
    db = sqlite3.connect(WAREHOUSE)
    db.execute('create table if not exists inventory_snapshot (as_of text, sku text, description text, category text, location text, '
               'qty integer, unit_cost real, ext_cost real, last_sale text)')
    db.executemany('insert into inventory_snapshot values (?,?,?,?,?,?,?,?,?)', [list(r.values()) for r in rows])
    db.commit()
    db.close()


def deliver_mail(label):
    for f in sorted((SRC / label / 'sources').glob('*.eml')):
        shutil.copy(f, MAILBOX / f.name)
        print(f'  mailbox got {f.name}')


def weekly_uploads(label):
    return [f for f in sorted((SRC / label / 'inputs').iterdir())]


def create(label, by='alex'):
    call('POST', '/' + label, by=by)


def close(label):
    call('POST', f'/{label}/close', by='alex')
    print(f'  closed {label}')


start = sys.argv[4] if len(sys.argv) > 4 else ORDER[0]
call('POST')
MAILBOX.mkdir(parents=True, exist_ok=True)
for label in ORDER:
    if label < start:
        continue
    print(f'\n== {label}')
    if label == '2026-08-16':
        for past in ORDER[:ORDER.index(label) + 1]:
            stage_sources(past)
        deliver_mail(label)
        create(label)
        for f in weekly_uploads(label):
            upload(label, f, 'in')
        steer(label, SOURCES_STEER)
        s = wait(label)
        screen(s)
        score(label, s)
        check(label, s)
        print('  08-16 stays open while 08-23 starts')
        continue
    if label == '2026-08-23':
        stage_sources(label)
        create(label)
        for f in weekly_uploads(label):
            upload(label, f, 'in')
        s = wait(label)
        screen(s)
        check(label, s)
        deliver_mail(label)
        steer(label, "Granite Point's email for this week has just arrived in the mailbox. Please fetch it again.")
        s = wait(label)
        screen(s)
        score(label, s)
        check(label, s)
        close('2026-08-16')
        close(label)
        create('scratch')
        call('POST', '/scratch/abandon', by='alex')
        print('  created and abandoned scratch')
        continue
    if label == '2026-08-30':
        stage_sources(label)
        deliver_mail(label)
        create(label)
        for f in weekly_uploads(label):
            upload(label, f, 'in')
        upload(label, next((SRC / label / 'extra').iterdir()), 'in', by='sam')
        s = wait(label)
        screen(s)
        score(label, s)
        check(label, s)
        close(label)
        continue
    create(label)
    files = weekly_uploads(label)
    if label == '2026-08-09':
        card = next(f for f in files if f.name.startswith('Card'))
        inv = next(f for f in files if f.name.startswith('Inventory'))
        for f in files:
            if f not in (card, inv):
                upload(label, f, 'in')
        upload(label, next((SRC / label / 'bad').iterdir()), 'in')
        started(label)
        time.sleep(90)
        upload(label, inv, 'in')
        s = wait(label)
        screen(s)
        upload(label, card, 'in')
        s = wait(label)
        screen(s)
        score(label, s)
        out = next(i for i in s['doc']['items'] if i['kind'] == 'output')
        upload(label, next((SRC / label / 'revision').iterdir()), 'out', by='sam', onto=out['id'])
        s = wait(label)
        screen(s)
        q = next((n for i in s['doc']['items'] for n in i.get('notes') or [] if n.get('options')
                  and not any(e['op'] == 'answer' and e.get('note') == n['id'] for e in s['log'])), None)
        if q:
            choice = next((k for k, o in enumerate(q['options']) if 'standing' in o.lower() or 'from now on' in o.lower()), 1)
            call('POST', f'/{label}/answer', json.dumps({'note': q['id'], 'choice': choice}).encode(), by='sam')
            print(f"  sam answered {q['id']}: {q['options'][choice]}")
            s = wait(label)
            screen(s)
        else:
            print('  NO QUESTION after the revision')
    else:
        for f in files:
            upload(label, f, 'in')
        if label in EXAMPLES:
            for f in sorted((SRC / label / 'outputs').iterdir()):
                upload(label, f, 'out')
        s = wait(label)
        screen(s)
        if label not in EXAMPLES:
            score(label, s)
    check(label, s)
    close(label)
