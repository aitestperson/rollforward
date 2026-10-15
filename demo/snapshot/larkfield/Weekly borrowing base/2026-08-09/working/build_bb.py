#!/usr/bin/env python3
"""Build the Larkfield weekly borrowing base from an instance's inputs.

Usage: build_bb.py INSTANCE_DIR [OUT_NAME]
Inputs are located through workspace.json item ids (ar-aging, card-receipts,
inventory, loan-activity); falls back to filename globs in upload/.
Writes working/<OUT_NAME> (default "Borrowing base Larkfield MM.DD.YYYY.xlsx")
and prints the computed lines as JSON.
"""
import csv, datetime as dt, glob, json, os, sys
import openpyxl

# Facility terms (Granite Point Credit), from the 2026-07-05 example output.
AR_ADVANCE = 0.80
CARD_ADVANCE = 0.90
INV_ADVANCE = 0.50
CONC_LIMIT = 0.25          # share of eligible AR (after >90 and cross-aged)
CROSS_AGE_SHARE = 0.50     # account ineligible if >50% of its balance is >90
SLOW_DAYS = 270            # no sale in more than 270 days before inventory date
AVAIL_BLOCK = 750000.00  # raised from 500,000 by Sam (reviewer) wk 08/09/2026, standing

GLOBS = {
    "ar-aging": "*AR aging*.xlsx",
    "card-receipts": "Card receipts*.csv",
    "inventory": "Inventory at cost*.xlsx",
    "loan-activity": "*loan activity*.xlsx",
}


# Inputs that may be missing for a week: built as 0 (conservative) and flagged.
OPTIONAL = {"card-receipts"}


def locate(inst):
    ws = json.load(open(os.path.join(inst, "workspace.json")))
    by_id = {i["id"]: i.get("file") for i in ws["items"]}
    out = {}
    for k, g in GLOBS.items():
        f = by_id.get(k)
        if not f:
            hits = sorted(glob.glob(os.path.join(inst, "upload", g)))
            if not hits:
                if k in OPTIONAL:
                    out[k] = None  # missing: line built as 0 and flagged
                    continue
                sys.exit(f"missing input {k}")
            f = os.path.relpath(hits[-1], inst)
        out[k] = os.path.join(inst, f)
    return out


def r2(x):
    # Round half up on the decimal value (strip float noise first): 0.545 -> 0.55,
    # matching the supplied workbooks (2026-07-19 had two half-cent lines).
    from decimal import Decimal, ROUND_HALF_UP
    return float(Decimal(repr(round(x + 0.0, 7))).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def compute(files, inv_date):
    # AR aging: Customer, Currency, Current, 1-30, 31-60, 61-90, 91-120, 120+, Total
    ws = openpyxl.load_workbook(files["ar-aging"], data_only=True).active
    hdr = [c.value for c in ws[1]]
    assert hdr[:9] == ["Customer", "Currency", "Current", "1-30", "31-60", "61-90", "91-120", "120+", "Total"], hdr
    ar = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        cust, cur, *b, tot = row[:9]
        assert cur == "USD", (cust, cur)
        assert abs(sum(b) - tot) < 0.01, cust
        ar.append(dict(cust=cust, total=tot, over90=b[4] + b[5]))
    gross = sum(a["total"] for a in ar)
    over90 = sum(a["over90"] for a in ar)
    cross = [a for a in ar if a["over90"] > CROSS_AGE_SHARE * a["total"]]
    cross_amt = sum(a["total"] - a["over90"] for a in cross)
    pre_conc = gross - over90 - cross_amt
    cap = CONC_LIMIT * pre_conc
    conc = []
    for a in ar:
        if a in cross:
            continue
        elig = a["total"] - a["over90"]
        if elig > cap:
            conc.append((a["cust"], elig - cap))
    conc_amt = sum(x for _, x in conc)

    # Card receipts in transit: sum of Net
    rows = list(csv.DictReader(open(files["card-receipts"], newline=""))) if files["card-receipts"] else []
    card = sum(float(r["Net"]) for r in rows)

    # Inventory
    ws = openpyxl.load_workbook(files["inventory"], data_only=True).active
    hdr = [c.value for c in ws[1]]
    assert hdr[:8] == ["SKU", "Description", "Category", "Location", "Qty", "Unit cost", "Ext cost", "Last sale"], hdr
    raw = [(i, r) for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2) if r[0]]
    # Exact duplicate rows (same SKU and all fields) are export repeats: keep the first.
    seen, inv, dup_rows = set(), [], []
    for i, r in raw:
        key = tuple(r[:8])
        if key in seen:
            dup_rows.append((i, r[0], r[6]))
            continue
        seen.add(key)
        inv.append(r)
    skus = [r[0] for r in inv]
    assert len(skus) == len(set(skus)), "same SKU with different data - resolve by hand"
    inv_total = sum(r[6] for r in inv)
    in_transit = sum(r[6] for r in inv if r[3] == "In-transit")
    slow = sum(r[6] for r in inv if r[3] != "In-transit" and (inv_date - r[7]).days > SLOW_DAYS)

    # Loan activity: closing balance
    ws = openpyxl.load_workbook(files["loan-activity"], data_only=True).active
    loan = {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    opening, closing = loan["Opening balance"], loan["Closing balance"]
    assert abs(opening + loan["Advances"] + loan["Repayments"] + loan["Interest"] - closing) < 0.01

    L = {}
    L["gross_ar"] = r2(gross)
    L["over90"] = -r2(over90)
    L["cross"] = -r2(cross_amt)
    L["conc"] = -r2(conc_amt)
    L["elig_ar"] = r2(L["gross_ar"] + L["over90"] + L["cross"] + L["conc"])
    L["ar_avail"] = r2(L["elig_ar"] * AR_ADVANCE)
    L["card"] = r2(card)
    L["card_avail"] = r2(L["card"] * CARD_ADVANCE)
    L["inv"] = r2(inv_total)
    L["in_transit"] = -r2(in_transit)
    L["slow"] = -r2(slow)
    L["elig_inv"] = r2(L["inv"] + L["in_transit"] + L["slow"])
    L["inv_avail"] = r2(L["elig_inv"] * INV_ADVANCE)
    L["gross_avail"] = r2(L["ar_avail"] + L["card_avail"] + L["inv_avail"])
    L["block"] = -AVAIL_BLOCK
    L["bb"] = r2(L["gross_avail"] + L["block"])
    L["revolver"] = -r2(closing)
    L["excess"] = r2(L["bb"] + L["revolver"])
    detail = dict(cross=[a["cust"] for a in cross], conc=conc, card_rows=len(rows),
                  card_dates=sorted({r["Date"] for r in rows}), card_missing=files["card-receipts"] is None,
                  skus=len(inv), inv_rows=len(raw), dup_rows=len(dup_rows),
                  dup_rows_first_last=(dup_rows[:1], dup_rows[-1:]),
                  dup_ext=round(sum(x[2] for x in dup_rows), 2))
    return L, detail


LAYOUT = [
    ("Gross accounts receivable", "gross_ar"),
    ("Less: over 90 days past due", "over90"),
    ("Less: cross-aged accounts (over 50% past 90 days)", "cross"),
    ("Less: concentration over 25%", "conc"),
    ("Eligible accounts receivable", "elig_ar"),
    ("Advance rate", AR_ADVANCE),
    ("Accounts receivable availability", "ar_avail"),
    ("Card receipts in transit", "card"),
    ("Advance rate", CARD_ADVANCE),
    ("Card availability", "card_avail"),
    ("Inventory at cost", "inv"),
    ("Less: in-transit", "in_transit"),
    ("Less: slow-moving (no sale in 270 days)", "slow"),
    ("Eligible inventory", "elig_inv"),
    ("Advance rate", INV_ADVANCE),
    ("Inventory availability", "inv_avail"),
    ("Gross availability", "gross_avail"),
    ("Less: availability block", "block"),
    ("Borrowing base", "bb"),
    ("Less: revolver balance", "revolver"),
    ("Excess availability", "excess"),
]


def write(L, week_end, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Base"
    ws["B1"] = "Larkfield Housewares, Inc. - Borrowing base"
    ws["B2"] = f"Week ending {week_end:%m/%d/%Y}"
    ws["B3"] = "Lender: Granite Point Credit"
    for i, (label, key) in enumerate(LAYOUT):
        r = 5 + i
        ws.cell(r, 2, label)
        if isinstance(key, float):
            c = ws.cell(r, 3, key)
            c.number_format = "0.00%"
        else:
            c = ws.cell(r, 3, L[key])
            c.number_format = "#,##0.00;(#,##0.00)"
    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 18
    wb.save(path)


def main():
    inst = sys.argv[1].rstrip("/")
    week_end = dt.datetime.strptime(os.path.basename(os.path.abspath(inst)), "%Y-%m-%d")
    files = locate(inst)
    # Inventory date: from the file name (MM.DD.YYYY), else day before week end
    import re
    m = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", os.path.basename(files["inventory"]))
    inv_date = dt.datetime(int(m[3]), int(m[1]), int(m[2])) if m else week_end - dt.timedelta(days=1)
    L, detail = compute(files, inv_date)
    name = sys.argv[2] if len(sys.argv) > 2 else f"Borrowing base Larkfield {week_end:%m.%d.%Y}.xlsx"
    out = os.path.join(inst, "working", name)
    if os.path.exists(out):
        sys.exit(f"{out} exists (frozen?) - pass a new name")
    write(L, week_end, out)
    print(json.dumps(dict(out=out, inv_date=f"{inv_date:%Y-%m-%d}", lines=L, detail=detail), indent=1, default=str))


if __name__ == "__main__":
    main()
