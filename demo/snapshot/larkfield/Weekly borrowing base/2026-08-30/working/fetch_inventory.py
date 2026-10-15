#!/usr/bin/env python3
"""Fetch the inventory snapshot from the warehouse database (read-only).

Usage: fetch_inventory.py INSTANCE_DIR [DB]
Takes the latest inventory_snapshot.as_of on or before the week end (instance label)
and writes working/Inventory at cost MM.DD.YYYY.xlsx (as_of date) in the layout of the
formerly uploaded files: sheet Detail, SKU..Last sale. Prints as_of, rows, total."""
import datetime as dt, os, sqlite3, sys
import openpyxl

DB = "/data-sources/warehouse.sqlite"
inst = sys.argv[1].rstrip("/")
db = sys.argv[2] if len(sys.argv) > 2 else DB
week_end = os.path.basename(os.path.abspath(inst))
con = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
as_of = con.execute("select max(as_of) from inventory_snapshot where as_of <= ?", (week_end,)).fetchone()[0]
if not as_of:
    sys.exit(f"no snapshot on or before {week_end}")
rows = con.execute("""select sku, description, category, location, qty, unit_cost, ext_cost, last_sale
                      from inventory_snapshot where as_of = ? order by sku""", (as_of,)).fetchall()
d = dt.date.fromisoformat(as_of)
out = os.path.join(inst, "working", f"Inventory at cost {d:%m.%d.%Y}.xlsx")
if os.path.exists(out):
    sys.exit(f"{out} exists (frozen?)")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Detail"
ws.append(["SKU", "Description", "Category", "Location", "Qty", "Unit cost", "Ext cost", "Last sale"])
for r in rows:
    r = list(r)
    r[7] = dt.datetime.fromisoformat(r[7])
    ws.append(r)
    ws.cell(ws.max_row, 6).number_format = "#,##0.00"
    ws.cell(ws.max_row, 7).number_format = "#,##0.00"
    ws.cell(ws.max_row, 8).number_format = "yyyy-mm-dd"
wb.save(out)
print(dict(out=out, db=db, as_of=as_of, rows=len(rows), skus=len({r[0] for r in rows}),
           ext_cost=round(sum(r[6] for r in rows), 2),
           fetched=dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%MZ")))
